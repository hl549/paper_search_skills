#!/usr/bin/env python3
"""Export a filtered-papers JSON (filter_papers.py output) to an Excel .xlsx
in the standard 'Articles' sheet layout: No. / Title / Journal / Year / Web Link(DOI).

Usage:
  python3 papers_to_excel.py --filtered electro_filtered.json \
      --out articles.xlsx [--title "Sheet name"] [--top N]
Requires openpyxl (use .venv/bin/python in the paper-digest dir, or `uv venv`).
"""
import argparse, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


HDR_FILL = PatternFill("solid", fgColor="1F4E79")
WRAP = Alignment(wrap_text=True, vertical="top")


def year_of(p):
    d = p.get("date") or ""
    if isinstance(d, str) and d[:4].isdigit():
        return d[:4]
    return p.get("year") or "-"


def link_of(p):
    doi = (p.get("doi") or "").strip()
    if doi:
        return f"https://doi.org/{doi}"
    return p.get("url") or ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--filtered", required=True, help="filter_papers.py output json (has 'kept')")
    ap.add_argument("--out", required=True)
    ap.add_argument("--title", default="Articles")
    ap.add_argument("--top", type=int, default=0, help="only top-N by score (0 = all kept)")
    args = ap.parse_args()

    data = json.load(open(args.filtered))
    papers = data.get("kept", data if isinstance(data, list) else [])
    papers = sorted(papers, key=lambda p: (-p.get("score", 0), p.get("date") or ""))
    if args.top > 0:
        papers = papers[:args.top]

    wb = Workbook()
    ws = wb.active
    ws.title = args.title
    headers = ["No.", "Title", "Journal", "Year", "Web Link (DOI)"]
    ws.append(headers)
    for c in ws[1]:
        c.fill = HDR_FILL
        c.font = Font(bold=True, color="FFFFFF")

    for i, p in enumerate(papers, 1):
        ws.append([i, p.get("title", ""), p.get("journal", ""), year_of(p), link_of(p)])
        ws.cell(row=i + 1, column=2).alignment = WRAP
        ws.cell(row=i + 1, column=5).alignment = WRAP

    for col, w in zip("ABCDE", [6, 70, 34, 8, 44]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    wb.save(args.out)
    print(f"SAVED {args.out} rows={len(papers)}")


if __name__ == "__main__":
    main()
